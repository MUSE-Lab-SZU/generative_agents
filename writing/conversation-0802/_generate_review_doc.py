#!/usr/bin/env python3
"""
生成面向心理老师的咨询对话审阅 Word 文档（v3）。

变更：
- 去除 KBD2-G4，KBD2 与其他案例一致
- 每个 KBD 仅保留质量最优的一组对话（-01 或 -02）
- 标题层级：H1=案例，H2=对话记录/评估表/附录案例，H3=单次咨询
- TOC 仅显示 H1-H2，单次咨询标题不出现在目录
- 对话导航网格使用 bookmark 超链接
"""

import json
import os
import re
import glob
from collections import defaultdict
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

# ============================================================
# Globals
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DOC = os.path.join(BASE_DIR, "AI_Counseling_Review_v3.docx")

# Anonymization mapping
PATIENT_MAP = {
    "KBD1": "案例 A",
    "KBD2": "案例 B",
    "KBD3": "案例 C",
    "KBD5": "案例 D",
    "KBD6": "案例 E",
    "KBD7": "案例 F",
}

# Select best run per KBD (based on quality analysis)
# KBD1: both runs similar quality, use -01
# KBD2: -02 much richer (18.4 vs 9.3 avg exchanges)
# KBD3: -02 much richer (14.4 vs 9.8, no short sessions)
# KBD5: -01 has patient avg 52 chars/msg (poor engagement), use -02
# KBD6: both runs similar, use -01
# KBD7: both runs similar, use -01
SELECTED_RUN = {
    "KBD1": "01",
    "KBD2": "02",
    "KBD3": "02",
    "KBD5": "02",
    "KBD6": "01",
    "KBD7": "01",
}

# Bookmark ID counter
_bm_counter = [0]

def _next_bm_id():
    _bm_counter[0] += 1
    return str(_bm_counter[0])

# ============================================================
# Bookmark & Hyperlink helpers (python-docx internal)
# ============================================================

def add_bookmark(paragraph, bookmark_name):
    """Add a Word bookmark to the end of a paragraph's runs."""
    bm_id = _next_bm_id()

    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), bm_id)
    start.set(qn('w:name'), bookmark_name)
    paragraph._element.append(start)

    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), bm_id)
    paragraph._element.append(end)


def add_internal_hyperlink(paragraph, text, bookmark_name, font_size=10,
                          color='0563C1', bold=False):
    """Add a clickable internal hyperlink to a paragraph."""
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), bookmark_name)
    hyperlink.set(qn('w:history'), '1')

    run_elem = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    # Font size
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(font_size * 2))  # half-points
    rPr.append(sz)

    # Color (blue for links)
    c = OxmlElement('w:color')
    c.set(qn('w:val'), color)
    rPr.append(c)

    # Underline
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)

    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)

    run_elem.append(rPr)

    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run_elem.append(t)
    hyperlink.append(run_elem)
    paragraph._element.append(hyperlink)

    return hyperlink


# ============================================================
# Heading helpers
# ============================================================

def add_heading_bm(doc, text, level, bookmark_name=None):
    """Add a Word heading with optional bookmark for navigation."""
    h = doc.add_heading(text, level=level)
    if bookmark_name:
        add_bookmark(h, bookmark_name)
    return h


def add_para(doc, text, bold=False, size=10.5, alignment=None,
             space_after=6, space_before=0, color=None, indent_left=None):
    """Add a formatted paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    if bold:
        run.bold = True
    if color:
        run.font.color.rgb = color
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent_left:
        p.paragraph_format.left_indent = Cm(indent_left)
    return p


# ============================================================
# TOC field
# ============================================================

def insert_toc_field(doc):
    """Insert a Word TOC field that auto-generates from headings."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Begin field
    run = p.add_run()
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')
    run._element.append(fldChar_begin)

    # Instruction
    run2 = p.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' TOC \\o "1-2" \\h \\z \\u '
    run2._element.append(instrText)

    # Separate
    run3 = p.add_run()
    fldChar_sep = OxmlElement('w:fldChar')
    fldChar_sep.set(qn('w:fldCharType'), 'separate')
    run3._element.append(fldChar_sep)

    # Placeholder text
    run4 = p.add_run('[ 请在 Word 中右键此处 → 更新域，以自动生成目录 ]')
    run4.font.size = Pt(9)
    run4.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    # End field
    run5 = p.add_run()
    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')
    run5._element.append(fldChar_end)

    doc.add_paragraph()  # spacer


# ============================================================
# Data extraction
# ============================================================

def extract_counseling_sessions(json_path):
    """Extract counseling sessions from JSON. Returns list of dicts."""
    with open(json_path, encoding='utf-8') as f:
        data = json.load(f)

    sessions = []
    for date_key in sorted(data.keys()):
        for sess_idx, session_group in enumerate(data[date_key]):
            for conv_key, messages in session_group.items():
                if 'persona' in conv_key:
                    conv = [(msg[0], msg[1]) for msg in messages]
                    sessions.append({
                        'date': date_key,
                        'session_idx': sess_idx,
                        'conv_key': conv_key,
                        'messages': conv,
                    })
    return sessions


def parse_filename(filename):
    """Parse filename to extract patient, group, run info."""
    basename = os.path.splitext(filename)[0]
    m = re.match(r'(?i)(kbd\d+)-g(\d+)-(\d+)-?(?:sev-)?conversation', basename)
    if m:
        return {
            'patient': m.group(1).upper(),
            'group': m.group(2),
            'run': m.group(3),
        }
    return None


# ============================================================
# Summary MD parsing (for appendix)
# ============================================================

def parse_summary_md(filepath):
    """Parse summary MD and extract scores. Returns dict."""
    with open(filepath, encoding='utf-8') as f:
        content = f.read()

    result = {
        'condition': '', 'variant': '', 'group': '', 'severity': '',
        'run_name': '', 'phq9': [], 'bdi2': [],
    }

    m = re.search(r'- variant:\s*`(\w+)`', content)
    if m:
        result['variant'] = m.group(1)
    m = re.search(r'- group:\s*`(\w+)`', content)
    if m:
        result['group'] = m.group(1).replace('g', '')
    m = re.search(r'- severity:\s*`(\w+)`', content)
    if m:
        result['severity'] = m.group(1)
    m = re.search(r'- run_name:\s*`([^`]+)`', content)
    if m:
        result['run_name'] = m.group(1)

    # Parse PHQ-9
    phq9_sec = re.search(r'#### PHQ-9\n\n(.*?)(?:\n####|\Z)', content, re.DOTALL)
    if phq9_sec:
        for line in phq9_sec.group(1).split('\n'):
            m2 = re.match(
                r'\|\s*(\S+)\s*\|\s*\d+\s*\|\s*\S+\s*\|\s*([\d.]+)\s*\|\s*(.+?)\s*\|',
                line
            )
            if m2:
                result['phq9'].append({
                    'point': m2.group(1),
                    'score': float(m2.group(2)),
                    'severity': m2.group(3).strip(),
                })

    # Parse BDI-II
    bdi2_sec = re.search(r'#### BDI-II\n\n(.*?)(?:\n##|\Z)', content, re.DOTALL)
    if bdi2_sec:
        for line in bdi2_sec.group(1).split('\n'):
            m2 = re.match(
                r'\|\s*(\S+)\s*\|\s*\d+\s*\|\s*\S+\s*\|\s*([\d.]+)\s*\|\s*(.+?)\s*\|',
                line
            )
            if m2:
                result['bdi2'].append({
                    'point': m2.group(1),
                    'score': float(m2.group(2)),
                    'severity': m2.group(3).strip(),
                })

    return result


def load_summaries():
    """Load all summary data, keyed by variant-group-run."""
    summaries = {}
    for f in sorted(glob.glob('*_summary.md')):
        data = parse_summary_md(f)
        v = data['variant']
        g = data['group']
        r = '01' if '-0802-01_' in f else '02'
        key = f"{v}-g{g}-{r}"
        summaries[key] = data
    return summaries


# ============================================================
# Session navigation grid
# ============================================================

def add_session_nav_grid(doc, session_labels, bookmark_prefix):
    """
    Add a navigation table with links to each session.
    session_labels: list of (label, bookmark_suffix) tuples.
    """
    n = len(session_labels)
    if n == 0:
        return

    # Calculate grid: prefer 5 columns
    cols = min(5, n)
    rows = (n + cols - 1) // cols

    add_para(doc, "对话导航：", bold=True, size=10, space_after=4, space_before=12)

    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i in range(rows * cols):
        r = i // cols
        c = i % cols
        cell = table.rows[r].cells[c]

        if i < n:
            label, bm_suffix = session_labels[i]
            bm_name = f"{bookmark_prefix}_{bm_suffix}"
            # Clear default paragraph and add hyperlink
            p = cell.paragraphs[0]
            p.clear()
            add_internal_hyperlink(p, label, bm_name, font_size=9, bold=False)
        else:
            p = cell.paragraphs[0]
            p.clear()
            run = p.add_run('')
            run.font.size = Pt(9)

    doc.add_paragraph()  # spacer


# ============================================================
# Dialogue block
# ============================================================

def add_dialogue_block(doc, session_num, total_sessions, date_str, messages,
                       bookmark_name):
    """Add a formatted counseling session dialogue block with bookmark."""
    # Session heading
    h = doc.add_heading(f"第 {session_num} 次咨询  |  {date_str}", level=3)
    add_bookmark(h, bookmark_name)

    # Separator
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pBdr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single',
        qn('w:sz'): '4',
        qn('w:space'): '1',
        qn('w:color'): 'CCCCCC',
    })
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Dialogue messages
    for speaker, text in messages:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.5)

        if '蜻蜓' in speaker:
            label = '咨询师'
            color = RGBColor(0x2E, 0x7D, 0x32)
        else:
            label = '来访者'
            color = RGBColor(0x37, 0x47, 0x4F)

        run_l = p.add_run(f"【{label}】")
        run_l.bold = True
        run_l.font.size = Pt(9.5)
        run_l.font.color.rgb = color

        run_t = p.add_run(f" {text}")
        run_t.font.size = Pt(9.5)

    # End spacer
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ============================================================
# Evaluation form (revised)
# ============================================================

def add_evaluation_form(doc, case_label, case_bm_prefix):
    """Add comprehensive evaluation form."""
    doc.add_page_break()

    add_heading_bm(doc, f"【{case_label}】专业评估表", level=2,
                   bookmark_name=f"{case_bm_prefix}_eval")

    add_para(doc,
        "请根据上述对话内容，从专业角度对这位来访者进行评估。"
        "请勿翻阅文档末尾的附录（系统评估数据），仅基于对话内容做出独立判断。",
        size=9.5, space_after=12)

    # ---- Part 1: 对话模拟质量 ----
    add_heading_bm(doc, "一、对话模拟质量评估", level=3,
                   bookmark_name=f"{case_bm_prefix}_eval_p1")

    q1_fields = [
        ("来访者的语言表达、情绪反应和思维模式\n是否符合抑郁症患者的临床特点？",
         "□ 非常符合  □ 比较符合  □ 部分符合  □ 不太符合  □ 完全不符合"),
        ("具体哪些表现符合/不符合抑郁特征？（请举例）", ""),
        ("咨询师的回应是否具备专业心理咨询师\n应有的水平（共情、引导、技术运用等）？",
         "□ 非常专业  □ 比较专业  □ 一般  □ 不够专业  □ 很不专业"),
        ("咨询师的专业性具体体现在哪些方面？\n不足之处有哪些？", ""),
        ("整体对话的自然度和真实感如何？\n是否感觉像真实的人类对话？",
         "□ 非常真实  □ 比较真实  □ 一般  □ 不够真实  □ 很不真实"),
    ]

    for i, (q, placeholder) in enumerate(q1_fields):
        add_para(doc, f"{i+1}. {q}", bold=True, size=10, space_after=2, space_before=6)
        if placeholder:
            add_para(doc, f"   答：{placeholder}", size=10, space_after=4, color=RGBColor(0x88, 0x88, 0x88))
        else:
            add_para(doc, "   答：", size=10, space_after=4)
            for _ in range(3):
                add_para(doc, "   " + "_" * 80, size=9, space_after=1,
                        color=RGBColor(0xCC, 0xCC, 0xCC))

    # ---- Part 2: 抑郁严重程度 ----
    add_heading_bm(doc, "二、抑郁严重程度判断", level=3,
                   bookmark_name=f"{case_bm_prefix}_eval_p2")

    sev_table = doc.add_table(rows=8, cols=3)
    sev_table.style = 'Table Grid'
    sev_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    sev_fields = [
        ("评估维度", "您的判断", "备注/依据"),
        ("整体抑郁程度", "□ 无  □ 轻度  □ 中度  □ 中重度  □ 重度", ""),
        ("PHQ-9 估算范围", "______ 至 ______ 分（满分27）", ""),
        ("BDI-II 估算范围", "______ 至 ______ 分（满分63）", ""),
        ("病情变化趋势", "□ 明显改善  □ 轻度改善  □ 无变化  □ 轻度恶化  □ 明显恶化", ""),
        ("如改善，从第几次咨询开始？", "第 ______ 次", ""),
        ("自伤/自杀风险", "□ 无风险  □ 被动念头  □ 主动念头  □ 具体计划  □ 无法判断", ""),
        ("总体印象与依据", "（请简要说明）", ""),
    ]

    for i, (dim, judge, note) in enumerate(sev_fields):
        row = sev_table.rows[i]
        for j, text in enumerate([dim, judge, note]):
            cell = row.cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
                set_cell_bg(cell, 'E8EAF6')

    doc.add_paragraph()

    # ---- Part 3: 核心症状 ----
    add_heading_bm(doc, "三、核心症状观察", level=3,
                   bookmark_name=f"{case_bm_prefix}_eval_p3")

    symp_table = doc.add_table(rows=11, cols=3)
    symp_table.style = 'Table Grid'
    symp_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    symp_fields = [
        ("症状维度", "是否观察到", "具体表现 / 对话线索"),
        ("情绪低落 / 悲伤", "□ 是  □ 否", ""),
        ("兴趣减退 / 快感缺失", "□ 是  □ 否", ""),
        ("精力减退 / 疲劳感", "□ 是  □ 否", ""),
        ("自我价值感低 / 无价值感", "□ 是  □ 否", ""),
        ("内疚感 / 自责", "□ 是  □ 否", ""),
        ("注意力/决策困难", "□ 是  □ 否", ""),
        ("精神运动性改变", "□ 是  □ 否", ""),
        ("睡眠问题", "□ 是  □ 否", ""),
        ("食欲/体重变化", "□ 是  □ 否", ""),
        ("消极思维/认知扭曲模式", "□ 是  □ 否", ""),
    ]

    for i, (dim, check, desc) in enumerate(symp_fields):
        for j, text in enumerate([dim, check, desc]):
            cell = symp_table.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
                set_cell_bg(cell, 'E8EAF6')

    doc.add_paragraph()

    # ---- Part 4: 治疗效果评估 ----
    add_heading_bm(doc, "四、治疗效果与临床合理性评估", level=3,
                   bookmark_name=f"{case_bm_prefix}_eval_p4")

    eff_fields = [
        ("从对话中是否能观察到治疗效果？\n（如认知重构、行为激活、情绪改善等）",
         "□ 效果明显  □ 有所改善  □ 改善不明显  □ 未见效果  □ 有所恶化"),
        ("治疗过程中CBT等技术的运用是否恰当？\n是否符合临床实践特点？", ""),
        ("来访者的变化速度是否符合临床预期？\n（考虑到咨询频率约为每3天一次）",
         "□ 变化速度合理  □ 变化偏快  □ 变化偏慢  □ 波动过大不合常理"),
        ("如果有不符合临床特点的地方，\n请具体指出：", ""),
    ]

    for i, (q, placeholder) in enumerate(eff_fields):
        add_para(doc, f"{i+1}. {q}", bold=True, size=10, space_after=2, space_before=6)
        if placeholder:
            add_para(doc, f"   答：{placeholder}", size=10, space_after=4,
                    color=RGBColor(0x88, 0x88, 0x88))
        else:
            add_para(doc, "   答：", size=10, space_after=4)
            for _ in range(3):
                add_para(doc, "   " + "_" * 80, size=9, space_after=1,
                        color=RGBColor(0xCC, 0xCC, 0xCC))

    # ---- Part 5: 量表评分合理性 ----
    add_heading_bm(doc, "五、量表评分合理性判断", level=3,
                   bookmark_name=f"{case_bm_prefix}_eval_p5")

    scale_fields = [
        ("根据对话内容，您认为该来访者在PHQ-9量表中的\n各条目应如何评分？（可逐条评估或给出总分范围）", ""),
        ("根据对话内容，该来访者在BDI-II量表中的\n评分应如何？", ""),
        ("您认为LLM自动评分的PHQ-9/BDI-II量表分数\n是否可能客观反映对话中表现的症状严重程度？",
         "□ 能客观反映  □ 基本能反映  □ 部分能反映\n□ 可能高估  □ 可能低估  □ 无法判断"),
        ("如果您认为量表评分与对话表现可能不一致，\n主要原因是什么？", ""),
    ]

    for i, (q, placeholder) in enumerate(scale_fields):
        add_para(doc, f"{i+1}. {q}", bold=True, size=10, space_after=2, space_before=6)
        if placeholder:
            add_para(doc, f"   答：{placeholder}", size=10, space_after=4,
                    color=RGBColor(0x88, 0x88, 0x88))
        else:
            add_para(doc, "   答：", size=10, space_after=4)
            for _ in range(3):
                add_para(doc, "   " + "_" * 80, size=9, space_after=1,
                        color=RGBColor(0xCC, 0xCC, 0xCC))

    # ---- Part 6: 其他意见 ----
    add_heading_bm(doc, "六、其他专业意见（可选）", level=3,
                   bookmark_name=f"{case_bm_prefix}_eval_p6")

    for _ in range(6):
        add_para(doc, "_" * 90, size=9, space_after=2,
                color=RGBColor(0xCC, 0xCC, 0xCC))


def set_cell_bg(cell, color):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color,
        qn('w:val'): 'clear',
    })
    shading.append(shd)


# ============================================================
# Appendix: System evaluation data per case
# ============================================================

def add_appendix(doc, selected_entries, summaries):
    """Add appendix with system scale scores and comparison framework."""
    doc.add_page_break()

    add_heading_bm(doc, "附录：系统评估结果（PHQ-9 / BDI-II 量表数据）", level=1,
                   bookmark_name="appendix")

    add_para(doc,
        "说明：以下为 AI 系统通过 LLM 自动评分的 PHQ-9 和 BDI-II 量表结果。"
        "每次评估使用 10 次固定重复取均值。"
        "此部分仅供研究团队在老师完成独立评估后进行对照参考，请勿提前展示给评估者。",
        size=9, space_after=12)

    case_to_variant = {v: k for k, v in PATIENT_MAP.items()}

    # Find matching summary for each selected entry
    case_order = [
        ("案例 A", "kbd1", "01"),
        ("案例 B", "kbd2", "02"),
        ("案例 C", "kbd3", "02"),
        ("案例 D", "kbd5", "02"),
        ("案例 E", "kbd6", "01"),
        ("案例 F", "kbd7", "01"),
    ]

    for case_label, variant, run in case_order:
        # Find matching summary (G1 only, matching run)
        summary_key = f"{variant}-g1-{run}"
        if summary_key not in summaries:
            # Try alternative format
            for k in summaries:
                if k.startswith(f"{variant}-g1"):
                    summary_key = k
                    break

        if summary_key not in summaries:
            add_para(doc, f"[WARN] No summary found for {case_label}", size=9)
            continue

        entry = summaries[summary_key]

        add_heading_bm(doc, f"{case_label} 系统评估数据", level=2,
                       bookmark_name=f"appendix_{case_label}_data")

        add_para(doc,
            f"数据来源：{variant.upper()} G1 Run-{run}  |  "
            f"预设严重程度：{entry['severity']}",
            bold=True, size=9, space_after=4, space_before=8)

        # ---- Combined score table ----
        # Determine max columns
        n_points = max(len(entry['phq9']), len(entry['bdi2']))
        points = [s['point'] for s in entry['phq9']]

        combined_t = doc.add_table(rows=4, cols=n_points + 1)
        combined_t.style = 'Table Grid'
        combined_t.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        hdr_cells = ["评估点"] + points
        for j, h in enumerate(hdr_cells):
            cell = combined_t.rows[0].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.font.size = Pt(8)
            run.bold = True
            set_cell_bg(cell, 'E8EAF6')

        # PHQ-9 row
        phq9_cells = ["PHQ-9"]
        for s in entry['phq9']:
            phq9_cells.append(f"{s['score']:.1f} ({s['severity']})")
        for j, text in enumerate(phq9_cells):
            cell = combined_t.rows[1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(7)

        # BDI-II row
        bdi2_cells = ["BDI-II"]
        for s in entry['bdi2']:
            bdi2_cells.append(f"{s['score']:.1f} ({s['severity']})")
        for j, text in enumerate(bdi2_cells):
            cell = combined_t.rows[2].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(7)

        # Delta row
        delta_cells = ["变化Δ"]
        if entry['phq9']:
            t0_p = entry['phq9'][0]['score']
            for s in entry['phq9']:
                d = s['score'] - t0_p
                delta_cells.append(f"{d:+.1f}")
        for j, text in enumerate(delta_cells):
            cell = combined_t.rows[3].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(7)

        # Summary
        t0_phq9 = entry['phq9'][0]['score'] if entry['phq9'] else 0
        last_phq9 = entry['phq9'][-1]['score'] if entry['phq9'] else 0
        delta_phq9 = last_phq9 - t0_phq9

        t0_bdi2 = entry['bdi2'][0]['score'] if entry['bdi2'] else 0
        last_bdi2 = entry['bdi2'][-1]['score'] if entry['bdi2'] else 0
        delta_bdi2 = last_bdi2 - t0_bdi2

        add_para(doc,
            f"PHQ-9: {t0_phq9:.1f} → {last_phq9:.1f} (delta={delta_phq9:+.1f})  |  "
            f"BDI-II: {t0_bdi2:.1f} → {last_bdi2:.1f} (delta={delta_bdi2:+.1f})",
            size=9, space_after=4, space_before=4)

        # Comparison table
        add_para(doc, "老师判断 vs 系统评分 对比记录：", bold=True, size=9,
                space_after=4, space_before=8)

        comp_t = doc.add_table(rows=2, cols=7)
        comp_t.style = 'Table Grid'
        comp_t.alignment = WD_TABLE_ALIGNMENT.CENTER

        comp_headers = ["评估维度", "老师判断", "PHQ-9(T0)",
                       "PHQ-9等级", "BDI-II(T0)", "BDI-II等级", "是否一致"]
        for j, h in enumerate(comp_headers):
            cell = comp_t.rows[0].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.font.size = Pt(7)
            run.bold = True
            set_cell_bg(cell, 'E8EAF6')

        row_data = [
            "严重程度",
            "（待老师填写）",
            f"{t0_phq9:.1f}",
            entry['phq9'][0]['severity'] if entry['phq9'] else "N/A",
            f"{t0_bdi2:.1f}",
            entry['bdi2'][0]['severity'] if entry['bdi2'] else "N/A",
            "（待对比）",
        ]
        for j, text in enumerate(row_data):
            cell = comp_t.rows[1].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(7)

        doc.add_paragraph()

    # Overall comparison methodology
    add_heading_bm(doc, "对比分析方法说明", level=2,
                   bookmark_name="appendix_method")

    add_para(doc, "PHQ-9 严重程度分类标准：", bold=True, size=10, space_after=4)
    phq9_std = doc.add_table(rows=6, cols=2)
    phq9_std.style = 'Table Grid'
    for i, (sev, rng) in enumerate([
        ("无抑郁", "0-4"), ("轻度抑郁", "5-9"), ("中度抑郁", "10-14"),
        ("中重度抑郁", "15-19"), ("重度抑郁", "20-27"),
        ("注意", "若老师判断等级与系统等级相差不超过一级，视为\"相邻匹配\""),
    ]):
        for j, text in enumerate([sev, rng]):
            cell = phq9_std.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
                set_cell_bg(cell, 'E8EAF6')

    doc.add_paragraph()

    add_para(doc, "BDI-II 严重程度分类标准：", bold=True, size=10, space_after=4)
    bdi2_std = doc.add_table(rows=5, cols=2)
    bdi2_std.style = 'Table Grid'
    for i, (sev, rng) in enumerate([
        ("无抑郁", "0-13"), ("轻度抑郁", "14-19"), ("中度抑郁", "20-28"),
        ("重度抑郁", "29-63"),
        ("注意", "BDI-II 重度范围较宽(29-63)，需特别关注分数绝对值"),
    ]):
        for j, text in enumerate([sev, rng]):
            cell = bdi2_std.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
                set_cell_bg(cell, 'E8EAF6')

    doc.add_paragraph()

    add_para(doc,
        "对比指标说明：\n"
        "1. 严重等级一致率：老师判断等级与量表判定等级一致的比例\n"
        "2. 相邻等级匹配率：相差不超过一级的比例（如\"中度\"vs\"中重度\"）\n"
        "3. 分数范围命中率：老师估算的分数范围是否包含了系统均值\n"
        "4. 改善趋势一致性：老师观察到的变化方向是否与量表分数Δ方向一致\n"
        "5. 系统偏差：老师是否倾向于高估或低估严重程度\n"
        "6. 对话质量与量表一致性：若老师认为对话不真实/不专业，则量表分数的参考价值需重新评估",
        size=9, space_after=8)


# ============================================================
# Case section generation
# ============================================================

def generate_case_section(doc, case_label, setting_desc, entry, summaries):
    """Generate a case section with single run of sessions."""
    doc.add_page_break()

    bm_prefix = case_label.replace(' ', '_').replace('-', '_')

    # Heading 1: Case title
    add_heading_bm(doc, case_label, level=1, bookmark_name=f"{bm_prefix}_top")

    add_para(doc, f"咨询设置：{setting_desc}", size=10, space_after=4)

    # --- Quick nav links within case ---
    add_para(doc, "本案例快速导航：", bold=True, size=10, space_after=2, space_before=6)
    nav_p = doc.add_paragraph()
    add_internal_hyperlink(nav_p, "[ 对话记录 ]", f"{bm_prefix}_dialogues",
                          font_size=10, bold=True)
    nav_p.add_run("  |  ").font.size = Pt(10)
    add_internal_hyperlink(nav_p, "[ 专业评估表 ]", f"{bm_prefix}_eval",
                          font_size=10, bold=True)
    doc.add_paragraph()

    # --- Dialogue section (Heading 2) ---
    add_heading_bm(doc, "咨询对话记录", level=2,
                   bookmark_name=f"{bm_prefix}_dialogues")

    sessions = entry['sessions']
    add_para(doc,
        f"共 {len(sessions)} 次咨询对话，"
        f"时间跨度：{sessions[0]['date'][:8]} 至 {sessions[-1]['date'][:8]}",
        size=9, space_after=6)

    # --- Session navigation grid ---
    session_labels = []
    for i, sess in enumerate(sessions):
        label = f"第{i+1}次"
        bm_suffix = f"sess{i+1}"
        session_labels.append((label, bm_suffix))

    add_session_nav_grid(doc, session_labels, bm_prefix)

    # --- Render each session (Heading 3 = not in TOC) ---
    for i, sess in enumerate(sessions):
        bm_name = f"{bm_prefix}_sess{i+1}"

        # Session heading with H3 level (does NOT appear in TOC with \o "1-2")
        h = doc.add_heading(f"第 {i+1} 次咨询  |  {sess['date']}", level=3)
        add_bookmark(h, bm_name)

        # Thin separator
        p = doc.add_paragraph()
        pPr = p._element.get_or_add_pPr()
        pBdr = pPr.makeelement(qn('w:pBdr'), {})
        bottom = pBdr.makeelement(qn('w:bottom'), {
            qn('w:val'): 'single',
            qn('w:sz'): '4',
            qn('w:space'): '1',
            qn('w:color'): 'CCCCCC',
        })
        pBdr.append(bottom)
        pPr.append(pBdr)

        # Dialogue messages
        for speaker, text in sess['messages']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(0.5)

            if '蜻蜓' in speaker:
                label = '咨询师'
                color = RGBColor(0x2E, 0x7D, 0x32)
            else:
                label = '来访者'
                color = RGBColor(0x37, 0x47, 0x4F)

            run_l = p.add_run(f"【{label}】")
            run_l.bold = True
            run_l.font.size = Pt(9.5)
            run_l.font.color.rgb = color

            run_t = p.add_run(f" {text}")
            run_t.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- Evaluation form (Heading 2 = appears in TOC) ---
    add_evaluation_form(doc, case_label, bm_prefix)


# ============================================================
# Front matter
# ============================================================

def add_front_matter(doc):
    """Add cover page, instructions, and TOC."""
    # Title
    doc.add_paragraph()
    doc.add_paragraph()

    add_para(doc, "AI 模拟心理咨询对话记录", bold=True, size=22,
             alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_para(doc, "—— 专业审阅材料 ——", size=14,
             alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)

    # Instructions heading
    add_heading_bm(doc, "使用说明", level=1, bookmark_name="instructions")

    instructions = [
        "1. 本文档包含由 AI 模拟生成的 6 位来访者与同一位 AI 咨询师的对话记录。",
        "2. 所有来访者已做匿名化处理（案例 A-F），严重程度标签已移除。请您仅基于对话内容进行专业判断。",
        "3. 对话中「咨询师」为 AI 模拟的专业心理咨询师（采用 CBT 取向），「来访者」为 AI 模拟的有心理困扰的个体。",
        "4. 每个案例包含若干次咨询对话（按时间顺序排列），页面顶部有导航网格可快速跳转到任意一次咨询。",
        "5. 每个案例末尾附有《专业评估表》，共六个部分。请逐项填写。",
        "6. 请勿在评估过程中翻阅文档末尾的附录（系统评估数据）——我们希望获得您纯粹基于对话内容的独立判断。",
        "7. 完成评估后，可将本文档交回研究团队，我们会将您的判断与系统量表结果进行对比分析。",
        "8. 提示：在 Word 中可使用「视图 → 导航窗格」查看文档结构，Ctrl+点击目录或导航网格中的链接可跳转。",
    ]

    for instr in instructions:
        add_para(doc, instr, size=10, space_after=3)

    # TOC
    doc.add_page_break()
    add_heading_bm(doc, "目  录", level=1, bookmark_name="toc")
    add_para(doc,
        "（请在 Word 中右键下方灰色文字 → 更新域，以自动生成带页码的目录）",
        size=9, space_after=8,
        color=RGBColor(0x99, 0x99, 0x99))
    insert_toc_field(doc)



# ============================================================
# Main
# ============================================================

def main():
    os.chdir(BASE_DIR)

    print("=" * 60)
    print("Generating Teacher Review Document v3")
    print("=" * 60)

    # ---- Step 1: Collect conversation data ----
    print("\n[1/3] Loading conversation data...")
    json_files = sorted(glob.glob('*.json'))

    all_entries = []
    for jf in json_files:
        info = parse_filename(jf)
        if info is None:
            print(f"  [SKIP] {jf}")
            continue
        # Skip G4 entirely
        if info['group'] == '4':
            print(f"  [SKIP] G4 excluded: {jf}")
            continue
        sessions = extract_counseling_sessions(jf)
        if not sessions:
            print(f"  [WARN] No counseling in {jf}")
            continue
        all_entries.append({
            'filename': jf,
            'patient': info['patient'],
            'group': info['group'],
            'run': info['run'],
            'sessions': sessions,
        })
        print(f"  [OK] {jf}: {len(sessions)} counseling sessions")

    # Select best run per patient
    selected_entries = {}
    for e in all_entries:
        patient = e['patient']
        best_run = SELECTED_RUN.get(patient, '01')
        if e['run'] == best_run:
            selected_entries[patient] = e
            print(f"  [SELECTED] {patient} -> run {best_run} ({e['filename']})")

    # ---- Step 2: Load summary data for appendix ----
    print("\n[2/3] Loading system evaluation data...")
    summaries = load_summaries()
    print(f"  Loaded {len(summaries)} summary records")

    # ---- Step 3: Generate document ----
    print("\n[3/3] Generating Word document...")

    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Default font
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # Modify heading styles for better visual
    for lvl in range(1, 4):
        h_style = doc.styles[f'Heading {lvl}']
        h_style.font.name = '微软雅黑'
        h_style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # Front matter
    add_front_matter(doc)

    # ---- Cases ----
    case_order = [
        ("KBD1", "案例 A"),
        ("KBD2", "案例 B"),
        ("KBD3", "案例 C"),
        ("KBD5", "案例 D"),
        ("KBD6", "案例 E"),
        ("KBD7", "案例 F"),
    ]

    for patient, case_label in case_order:
        if patient in selected_entries:
            generate_case_section(doc, case_label,
                                "AI 模拟心理咨询（CBT取向）",
                                selected_entries[patient], summaries)

    # ---- Appendix ----
    add_appendix(doc, selected_entries, summaries)

    # ---- Save ----
    doc.save(OUTPUT_DOC)
    print(f"\n{'='*60}")
    print(f"Document saved to: {OUTPUT_DOC}")
    print(f"{'='*60}")
    print("\nFeatures:")
    print("  - H1=Case, H2=Dialogues/Eval, H3=Sessions (not in TOC)")
    print("  - TOC field (\\o 1-2) - update in Word")
    print("  - Session navigation grids with bookmark hyperlinks")
    print("  - 6 cases, best run per KBD, G4 excluded")
    print("  - Revised evaluation form (6 parts)")
    print("  - Appendix with PHQ-9/BDI-II scores per case")


if __name__ == '__main__':
    main()
